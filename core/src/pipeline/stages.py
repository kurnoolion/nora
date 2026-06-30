"""Pipeline stage functions.

Each stage function takes a PipelineContext and returns a StageResult.
Imports are lazy within each function to tolerate missing dependencies.
"""

from __future__ import annotations

import hashlib
import json
import logging
import shutil
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING

from core.src.pipeline.cells import Cell

if TYPE_CHECKING:
    from core.src.pipeline.runner import PipelineContext

logger = logging.getLogger(__name__)


@dataclass
class StageResult:
    """Result from running a single pipeline stage."""

    stage: str
    status: str  # OK, WARN, FAIL, SKIP
    elapsed_seconds: float
    stats: dict = field(default_factory=dict)
    error_code: str | None = None
    error_message: str | None = None
    warnings: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return self.status in ("OK", "WARN")


def _fail(stage: str, code: str, message: str, elapsed: float = 0.0) -> StageResult:
    return StageResult(
        stage=stage, status="FAIL", elapsed_seconds=elapsed,
        error_code=code, error_message=message,
    )


# ---------------------------------------------------------------------------
# Stage 1: extract
# ---------------------------------------------------------------------------

def run_extract(ctx: PipelineContext) -> StageResult:
    """Extract documents into normalized IR."""
    t0 = time.time()
    stage = "extract"
    out_dir = ctx.stage_output(stage)
    out_dir.mkdir(parents=True, exist_ok=True)

    try:
        from core.src.extraction.registry import extract_document, infer_metadata_from_path, supported_extensions
    except ImportError as e:
        return _fail(stage, "EXT-E001", f"Import error: {e}", time.time() - t0)

    exts = supported_extensions()
    files = sorted(
        f for f in ctx.documents_dir.rglob("*")
        if f.is_file() and f.suffix.lower() in exts
    )

    if not files:
        return _fail(stage, "ENV-E002", f"No documents in {ctx.documents_dir}", time.time() - t0)

    stats = {"docs": 0, "skipped": 0, "blocks": 0, "tables": 0, "failed": 0}
    warnings: list[str] = []
    ir_paths: list[str] = []

    for f in files:
        # infer_metadata_from_path enforces the MMMYYYY cell convention
        # fail-loud (EXT-E004); a mis-named release dir aborts the stage here.
        metadata = infer_metadata_from_path(f)
        # D-DRAFT-8: honor the --mno/--release cell scope.
        if not ctx.cell_in_scope(metadata["mno"], metadata["release"]):
            continue
        # D-DRAFT-6: route each IR to its (mno, release) cell directory.
        cell_dir = ctx.stage_output(stage, Cell(metadata["mno"], metadata["release"]))
        out_path = cell_dir / f"{f.stem}_ir.json"
        # D-DRAFT-8 skip: reuse an IR that's newer than its source. `--force` overrides.
        if not ctx.force and out_path.exists() and out_path.stat().st_mtime >= f.stat().st_mtime:
            ir_paths.append(str(out_path))
            stats["skipped"] += 1
            continue
        try:
            ir = extract_document(
                f, mno=metadata["mno"], release=metadata["release"],
                doc_type=metadata["doc_type"],
            )
            # Path-derived plan (per-plan input dir), if any — the extractors
            # don't see the path, so stamp it here for the parser to consume.
            ir.plan = metadata.get("plan", "")
            cell_dir.mkdir(parents=True, exist_ok=True)
            ir.save_json(out_path)
            ir_paths.append(str(out_path))
            stats["docs"] += 1
            stats["blocks"] += ir.block_count
            tbl_count = sum(1 for b in ir.content_blocks if b.type.value == "table")
            stats["tables"] += tbl_count
            if ir.block_count < 10:
                warnings.append(f"EXT-W001: Low blocks ({ir.block_count}) in {f.name}")
        except Exception as e:
            stats["failed"] += 1
            warnings.append(f"EXT-E001: {f.name}: {e}")

    ctx.state["ir_paths"] = ir_paths
    elapsed = time.time() - t0
    status = "WARN" if stats["failed"] > 0 else "OK"
    return StageResult(stage=stage, status=status, elapsed_seconds=elapsed,
                       stats=stats, warnings=warnings)


# ---------------------------------------------------------------------------
# Stage 2: profile
# ---------------------------------------------------------------------------

def run_profile(ctx: PipelineContext) -> StageResult:
    """Resolve + materialize each cell's parse profile (D-DRAFT-7).

    Per cell, the profile is resolved from `<env_dir>/profiles.json` (or the
    `--profile` / corrections override applied to every cell), placeholder-
    substituted, and written to `out/profile/<mno>/<rel>/profile.json` for the
    parse stage. No auto-profiling — an uncovered cell fails loud (PIP-E003),
    because these corpora use hand-authored profiles.
    """
    t0 = time.time()
    stage = "profile"

    cells = ctx.input_cells()
    if not cells:
        return _fail(stage, "PIP-E002",
                     f"No input cells under {ctx.documents_dir} "
                     f"(expected input/<MNO>/<MMMYYYY>/)", time.time() - t0)

    # `--profile` (CLI) or corrections/profile.json act as a single-profile
    # override applied to every cell (single-MNO back-compat / wildcard).
    override = ctx.state.get("profile_path")
    if not override:
        correction = ctx.correction("profile.json")
        if correction:
            override = str(correction)

    try:
        from core.src.env.profile_bindings import load_profile_bindings
        from core.src.profiler.profile_substitute import load_substituted_profile
    except ImportError as e:
        return _fail(stage, "PRF-E001", f"Import error: {e}", time.time() - t0)

    # <env_dir>/out/profile -> <env_dir>
    env_dir = ctx.stage_output(stage).parent.parent
    bindings = load_profile_bindings(env_dir, override=override)

    # Coverage check — fail loud listing every uncovered cell at once.
    uncovered = bindings.uncovered(cells)
    if uncovered:
        miss = ", ".join(f"{m}/{r}" for m, r in uncovered)
        return _fail(stage, "PIP-E003",
                     f"No profile bound for: {miss}. Add bindings to "
                     f"<env_dir>/profiles.json (or pass --profile).",
                     time.time() - t0)

    stats = {"cells": 0}
    for cell in cells:
        profile_path = bindings.resolve(cell.mno, cell.release)
        if not Path(profile_path).exists():
            return _fail(stage, "PIP-E002",
                         f"Bound profile not found for {cell.mno}/{cell.release}: "
                         f"{profile_path}", time.time() - t0)
        # Substitute placeholders ONCE here; parse reads the materialized
        # profile raw (already substituted).
        profile = load_substituted_profile(Path(profile_path), env_dir=env_dir)
        cell_out = ctx.stage_output(stage, cell)
        cell_out.mkdir(parents=True, exist_ok=True)
        profile.save_json(cell_out / "profile.json")
        stats["cells"] += 1

    return StageResult(stage=stage, status="OK",
                       elapsed_seconds=time.time() - t0, stats=stats)


# ---------------------------------------------------------------------------
# Stage 3: parse
# ---------------------------------------------------------------------------

def run_parse(ctx: PipelineContext) -> StageResult:
    """Parse each cell's IRs into requirement trees (D-DRAFT-6/7).

    Per cell, loads the cell's materialized profile from
    `out/profile/<mno>/<rel>/profile.json` (already placeholder-substituted by
    the profile stage) and parses `out/extract/<mno>/<rel>/*_ir.json` into
    `out/parse/<mno>/<rel>/*_tree.json`.
    """
    t0 = time.time()
    stage = "parse"

    try:
        from core.src.models.document import DocumentIR
        from core.src.profiler.profile_schema import DocumentProfile
        from core.src.parser.structural_parser import GenericStructuralParser, RequirementTree
        from core.src.parser.user_annotations import apply_user_annotations
    except ImportError as e:
        return _fail(stage, "PRS-E001", f"Import error: {e}", time.time() - t0)

    cells = ctx.input_cells()
    if not cells:
        return _fail(stage, "PIP-E002",
                     f"No input cells under {ctx.documents_dir}", time.time() - t0)

    stats = {"docs": 0, "skipped": 0, "reqs": 0, "max_depth": 0, "toc": 0, "struck": 0, "cascade": 0, "revhist": 0, "defs": 0, "user_removes": 0, "toc_pair_misses": 0, "frontmatter": 0}
    tree_paths: list[str] = []
    warnings: list[str] = []
    # Per-doc parse summaries accumulated across ALL cells; emitted as the
    # consolidated <env_dir>/reports/parse_summary.json for the Parse Review UI.
    per_doc_summaries: list = []
    # ctx.stage_output("parse") == <env_dir>/out/parse → parent.parent == <env_dir>
    _env_dir = ctx.stage_output(stage).parent.parent
    log_dir = _env_dir / "reports" / "parse_log"
    annotations_dir = _env_dir / "annotations"

    for cell in cells:
        # Cell profile materialized (and substituted) by the profile stage.
        profile_path = ctx.stage_output("profile", cell) / "profile.json"
        if not profile_path.exists():
            warnings.append(f"PIP-E002: no profile for {cell.mno}/{cell.release} "
                            f"(run the profile stage first)")
            continue
        profile = DocumentProfile.load_json(profile_path)
        parser = GenericStructuralParser(profile)
        # D-DRAFT-8: fingerprint of the (already-substituted) cell profile —
        # a changed profile/mapping flips it and forces a re-parse.
        profile_fp = hashlib.sha256(profile_path.read_bytes()).hexdigest()[:16]

        ir_files = sorted(ctx.stage_output("extract", cell).glob("*_ir.json"))
        if not ir_files:
            warnings.append(f"PIP-E002: no IRs for {cell.mno}/{cell.release}")
            continue
        cell_out = ctx.stage_output(stage, cell)
        cell_out.mkdir(parents=True, exist_ok=True)

        for f in ir_files:
            out_path = cell_out / (f.stem.replace("_ir", "_tree") + ".json")
            # D-DRAFT-8 skip: reuse an up-to-date tree (newer than its IR) whose
            # stamped fingerprint matches the current profile. `--force` overrides.
            if not ctx.force and out_path.exists() and out_path.stat().st_mtime >= f.stat().st_mtime:
                try:
                    if RequirementTree.load_json(out_path).profile_fingerprint == profile_fp:
                        tree_paths.append(str(out_path))
                        stats["skipped"] += 1
                        continue
                except Exception:
                    pass  # unreadable/old tree — fall through and re-parse
            try:
                doc = DocumentIR.load_json(f)
                # D-061: apply user `remove` annotations to the IR before parse.
                doc_id = Path(doc.source_file).stem
                ann_path = annotations_dir / f"{doc_id}_annotations.json"
                n_removes = apply_user_annotations(doc, ann_path)
                if n_removes:
                    stats["user_removes"] += n_removes
                tree = parser.parse(doc)
                tree.profile_fingerprint = profile_fp  # D-DRAFT-8 skip key
                tree.save_json(out_path)
                tree_paths.append(str(out_path))
                if tree.parse_log is not None:
                    log_path = log_dir / f"{Path(doc.source_file).stem}_parse_log.json"
                    tree.parse_log.save_json(log_path)
                stats["docs"] += 1
                stats["reqs"] += len(tree.requirements)
                depth = max((r.section_number.count(".") for r in tree.requirements), default=0) + 1
                stats["max_depth"] = max(stats["max_depth"], depth)
                stats["toc"] += tree.parse_stats.toc_blocks_dropped
                stats["struck"] += tree.parse_stats.struck_blocks_dropped
                stats["cascade"] += tree.parse_stats.cascade_blocks_dropped
                stats["revhist"] += tree.parse_stats.revhist_blocks_dropped
                stats["defs"] += tree.parse_stats.defs_extracted
                stats["toc_pair_misses"] += tree.parse_stats.toc_pair_misses
                stats["frontmatter"] += tree.parse_stats.frontmatter_blocks_dropped
                if tree.parse_summary is not None:
                    per_doc_summaries.append(tree.parse_summary)
            except Exception as e:
                warnings.append(f"PRS-E001: {f.name}: {e}")

    # Consolidated parse summary — feeds the Parse Review Summary tab.
    if per_doc_summaries:
        try:
            from core.src.parser.parse_summary import build_corpus_summary
            from datetime import datetime, timezone
            corpus = build_corpus_summary(
                per_doc_summaries,
                generated_at=datetime.now(timezone.utc).isoformat(),
            )
            summary_path = _env_dir / "reports" / "parse_summary.json"
            corpus.save_json(summary_path)
            stats["docs_missing_revhist"] = corpus.docs_without_revhist
            stats["docs_missing_glossary"] = corpus.docs_without_glossary
        except Exception as e:
            warnings.append(f"PRS-W005: parse_summary write failed: {e}")

    ctx.state["tree_paths"] = tree_paths
    return StageResult(stage=stage, status="WARN" if warnings else "OK",
                       elapsed_seconds=time.time() - t0, stats=stats, warnings=warnings)


# ---------------------------------------------------------------------------
# Stage 4: resolve
# ---------------------------------------------------------------------------

def run_resolve(ctx: PipelineContext) -> StageResult:
    """Resolve cross-references across parsed trees."""
    t0 = time.time()
    stage = "resolve"
    out_dir = ctx.stage_output(stage)
    out_dir.mkdir(parents=True, exist_ok=True)

    try:
        from core.src.parser.structural_parser import RequirementTree
        from core.src.resolver.resolver import CrossReferenceResolver
    except ImportError as e:
        return _fail(stage, "RES-E001", f"Import error: {e}", time.time() - t0)

    cells = ctx.input_cells()
    if not cells:
        return _fail(stage, "PIP-E002",
                     f"No input cells under {ctx.documents_dir}", time.time() - t0)

    stats = {"internal": 0, "cross_plan": 0, "standards": 0, "unresolved": 0}
    n_trees = 0
    # D-DRAFT-10: resolve PER CELL — each cell's trees only. Cross-references
    # therefore never match across MNOs or releases (plan codes/numbers aren't
    # globally unique); cross-cell relations are the global graph's job.
    for cell in cells:
        cell_trees_dir = ctx.stage_output("parse", cell)
        tree_files = sorted(cell_trees_dir.glob("*_tree.json"))
        if not tree_files:
            continue
        n_trees += len(tree_files)
        trees = [RequirementTree.load_json(f) for f in tree_files]
        manifests = CrossReferenceResolver(trees).resolve_all()
        cell_out = ctx.stage_output(stage, cell)
        cell_out.mkdir(parents=True, exist_ok=True)
        for m in manifests:
            m.save_json(cell_out / f"{m.plan_id}_xrefs.json")
            s = m.summary
            stats["internal"] += s.resolved_internal
            stats["cross_plan"] += s.resolved_cross_plan
            stats["standards"] += s.resolved_standards
            stats["unresolved"] += s.broken_internal

    if n_trees == 0:
        return _fail(stage, "PIP-E002",
                     f"No tree files under {ctx.stage_output('parse')}", time.time() - t0)

    warnings: list[str] = []
    if stats["unresolved"] > 0:
        warnings.append(f"RES-W001: {stats['unresolved']} unresolved internal refs")

    return StageResult(stage=stage, status="WARN" if warnings else "OK",
                       elapsed_seconds=time.time() - t0, stats=stats, warnings=warnings)


# ---------------------------------------------------------------------------
# Stage 5: taxonomy
# ---------------------------------------------------------------------------

def _corpus_fingerprint(tree_files: list[Path], parse_dir: Path) -> str:
    """Hash of the contributing tree set (D-DRAFT-9).

    Captures additions, removals, and edits — each tree's path (relative to
    `parse_dir`, so it's cell-stable) plus a hash of its bytes. A change in
    any tree flips the fingerprint and busts the taxonomy cache.
    """
    h = hashlib.sha256()
    for f in sorted(tree_files):
        try:
            rel = f.relative_to(parse_dir).as_posix()
        except ValueError:
            rel = f.name
        h.update(rel.encode("utf-8"))
        h.update(hashlib.sha256(f.read_bytes()).digest())
    return h.hexdigest()[:16]


def run_taxonomy(ctx: PipelineContext) -> StageResult:
    """Extract feature taxonomy from parsed trees."""
    t0 = time.time()
    stage = "taxonomy"
    out_dir = ctx.stage_output(stage)
    out_dir.mkdir(parents=True, exist_ok=True)

    # Check for correction override
    correction = ctx.correction("taxonomy.json")
    if correction:
        shutil.copy2(correction, out_dir / "taxonomy.json")
        ctx.state["taxonomy_path"] = str(out_dir / "taxonomy.json")
        return StageResult(
            stage=stage, status="OK", elapsed_seconds=time.time() - t0,
            stats={"source": "correction", "path": str(correction)},
            warnings=["TAX-W002: Using correction file for taxonomy"],
        )

    try:
        from core.src.parser.structural_parser import RequirementTree
        from core.src.taxonomy.extractor import FeatureExtractor
        from core.src.taxonomy.consolidator import TaxonomyConsolidator
    except ImportError as e:
        return _fail(stage, "TAX-E001", f"Import error: {e}", time.time() - t0)

    parse_dir = ctx.stage_output("parse")
    # rglob: taxonomy is a GLOBAL stage — it derives ONE union feature set
    # over every cell's trees (D-DRAFT-9), so it reads nested per-cell trees.
    tree_files = sorted(parse_dir.rglob("*_tree.json"))
    if not tree_files:
        return _fail(stage, "PIP-E002", f"No tree files in {parse_dir}", time.time() - t0)

    # D-DRAFT-9: corpus-fingerprint cache. The taxonomy is LLM-derived,
    # expensive, and non-deterministic across runs — so re-derive the (global,
    # union) taxonomy only when the contributing tree set changed; otherwise
    # reuse the cached taxonomy.json. `--force` busts the cache.
    taxonomy_path = out_dir / "taxonomy.json"
    fp_path = out_dir / ".corpus_fingerprint"
    corpus_fp = _corpus_fingerprint(tree_files, parse_dir)
    if (not ctx.force and taxonomy_path.exists() and fp_path.exists()
            and fp_path.read_text(encoding="utf-8").strip() == corpus_fp):
        ctx.state["taxonomy_path"] = str(taxonomy_path)
        return StageResult(
            stage=stage, status="OK", elapsed_seconds=time.time() - t0,
            stats={"source": "cache", "trees": len(tree_files)},
        )

    # Create LLM provider (extraction runs at temperature=0 — see
    # FeatureExtractor — for reproducible feature mappings).
    llm = ctx.create_llm_provider()
    extractor = FeatureExtractor(llm)
    all_doc_features = []

    for f in tree_files:
        tree = RequirementTree.load_json(f)
        doc_features = extractor.extract(tree)
        doc_out = out_dir / f"{tree.plan_id}_features.json"
        doc_features.save_json(doc_out)
        all_doc_features.append(doc_features)

    consolidator = TaxonomyConsolidator()
    taxonomy = consolidator.consolidate(all_doc_features)
    taxonomy.save_json(taxonomy_path)
    # D-DRAFT-9: stamp the corpus fingerprint so an unchanged re-run hits cache.
    fp_path.write_text(corpus_fp, encoding="utf-8")

    ctx.state["taxonomy_path"] = str(taxonomy_path)
    stats = {"features": len(taxonomy.features), "docs": len(all_doc_features), "source": "derived"}
    return StageResult(stage=stage, status="OK", elapsed_seconds=time.time() - t0, stats=stats)


# ---------------------------------------------------------------------------
# Stage 6: standards
# ---------------------------------------------------------------------------

def run_standards(ctx: PipelineContext) -> StageResult:
    """Ingest referenced 3GPP standards."""
    t0 = time.time()
    stage = "standards"
    out_dir = ctx.stage_output(stage)
    out_dir.mkdir(parents=True, exist_ok=True)

    try:
        from core.src.standards.reference_collector import StandardsReferenceCollector
        from core.src.standards.spec_downloader import SpecDownloader
        from core.src.standards.spec_parser import SpecParser
        from core.src.standards.section_extractor import SectionExtractor
    except ImportError as e:
        return _fail(stage, "STD-E001", f"Import error: {e}", time.time() - t0)

    resolve_dir = ctx.stage_output("resolve")
    parse_dir = ctx.stage_output("parse")

    collector = StandardsReferenceCollector()
    try:
        index = collector.collect(manifest_dir=resolve_dir, trees_dir=parse_dir)
    except Exception as e:
        return _fail(stage, "STD-E001", f"Reference collection failed: {e}", time.time() - t0)

    index_path = out_dir / "reference_index.json"
    index.save_json(index_path)

    # Download + parse + extract
    downloader = SpecDownloader(
        cache_dir=out_dir, source=ctx.standards_source
    )
    spec_parser = SpecParser()
    extractor = SectionExtractor()

    stats = {"specs_found": len(index.specs), "downloaded": 0, "parsed": 0, "extracted": 0, "failed": 0}
    warnings: list[str] = []

    for spec_ref in index.specs:
        if spec_ref.release_num <= 0:
            continue
        label = f"TS {spec_ref.spec} Rel-{spec_ref.release_num}"
        try:
            doc_path = downloader.download(spec_ref.spec, spec_ref.release_num)
            if not doc_path:
                warnings.append(f"STD-W001: {label} not found")
                stats["failed"] += 1
                continue
            stats["downloaded"] += 1

            spec_dir = out_dir / f"TS_{spec_ref.spec}" / f"Rel-{spec_ref.release_num}"
            spec_doc = spec_parser.parse(doc_path)
            spec_doc.save_json(spec_dir / "spec_parsed.json")
            stats["parsed"] += 1

            result = extractor.extract(spec_doc, spec_ref.sections, source_plans=spec_ref.source_plans)
            result.save_json(spec_dir / "sections.json")
            stats["extracted"] += 1
        except Exception as e:
            warnings.append(f"STD-E002: {label}: {e}")
            stats["failed"] += 1

    return StageResult(stage=stage, status="WARN" if stats["failed"] > 0 else "OK",
                       elapsed_seconds=time.time() - t0, stats=stats, warnings=warnings)


# ---------------------------------------------------------------------------
# Stage 7: graph
# ---------------------------------------------------------------------------

def run_graph(ctx: PipelineContext) -> StageResult:
    """Build the knowledge graph."""
    t0 = time.time()
    stage = "graph"
    out_dir = ctx.stage_output(stage)
    out_dir.mkdir(parents=True, exist_ok=True)

    try:
        from core.src.graph.builder import KnowledgeGraphBuilder
    except ImportError as e:
        return _fail(stage, "GRF-E001", f"Import error: {e}", time.time() - t0)

    builder = KnowledgeGraphBuilder()
    taxonomy_path = Path(ctx.state.get("taxonomy_path", ctx.stage_output("taxonomy") / "taxonomy.json"))

    try:
        graph = builder.build(
            trees_dir=ctx.stage_output("parse"),
            manifests_dir=ctx.stage_output("resolve"),
            taxonomy_path=taxonomy_path,
            standards_dir=ctx.stage_output("standards"),
        )
        graph_stats = builder.compute_stats()
    except Exception as e:
        return _fail(stage, "GRF-E001", f"Graph build failed: {e}", time.time() - t0)

    # Save
    import networkx as nx
    graph_path = out_dir / "knowledge_graph.json"
    with open(graph_path, "w") as f:
        json.dump(nx.node_link_data(graph), f, indent=2)
    graph_stats.save_json(out_dir / "graph_stats.json")

    ctx.state["graph_path"] = str(graph_path)

    # Connected components
    cc = nx.number_weakly_connected_components(graph)

    stats = {
        "nodes": graph_stats.total_nodes,
        "edges": graph_stats.total_edges,
        "components": cc,
    }
    warnings: list[str] = []
    if cc > 1:
        warnings.append(f"GRF-W001: {cc} connected components")

    return StageResult(stage=stage, status="WARN" if warnings else "OK",
                       elapsed_seconds=time.time() - t0, stats=stats, warnings=warnings)


# ---------------------------------------------------------------------------
# Stage 8: vectorstore
# ---------------------------------------------------------------------------

def run_vectorstore(ctx: PipelineContext) -> StageResult:
    """Build one vector store PER CELL (D-DRAFT-11).

    Each `(mno, release)` cell gets its own ChromaDB at
    `out/vectorstore/<mno>/<rel>/`, built from that cell's trees only — so
    retrieval statistics stay isolated per cell. The (global) taxonomy supplies
    feature_ids across cells. Query-side cell routing + fusion is the follow-on
    (D-DRAFT-11 slice 2).
    """
    t0 = time.time()
    stage = "vectorstore"

    try:
        from core.src.vectorstore import make_embedder
        from core.src.vectorstore.config import VectorStoreConfig
        from core.src.vectorstore.builder import VectorStoreBuilder
        from core.src.vectorstore.store_chroma import ChromaDBStore
    except ImportError as e:
        return _fail(stage, "VEC-E001", f"Import error: {e}", time.time() - t0)

    cells = ctx.input_cells()
    if not cells:
        return _fail(stage, "PIP-E002",
                     f"No input cells under {ctx.documents_dir}", time.time() - t0)

    env_dir = ctx.documents_dir.parent
    # One embedder, reused across cells (it doesn't depend on persist_directory).
    base_config = VectorStoreConfig(
        embedding_provider=ctx.embedding_provider,
        embedding_model=ctx.embedding_model,
        skipped_headers_log=str(env_dir / "reports" / "marker_headers.csv"),
    )
    embedder = make_embedder(base_config)
    taxonomy_path = Path(ctx.state.get("taxonomy_path", ctx.stage_output("taxonomy") / "taxonomy.json"))

    total_chunks = 0
    model = ""
    for cell in cells:
        cell_out = ctx.stage_output(stage, cell)
        cell_out.mkdir(parents=True, exist_ok=True)
        config = VectorStoreConfig(
            persist_directory=str(cell_out),
            embedding_provider=ctx.embedding_provider,
            embedding_model=ctx.embedding_model,
            skipped_headers_log=str(env_dir / "reports" / "marker_headers.csv"),
        )
        store = ChromaDBStore(
            persist_directory=config.persist_directory,
            collection_name=config.collection_name,
            distance_metric=config.distance_metric,
        )
        builder = VectorStoreBuilder(embedder=embedder, store=store, config=config)
        try:
            build_stats = builder.build(
                trees_dir=ctx.stage_output("parse", cell),  # this cell's trees only
                taxonomy_path=taxonomy_path,
                rebuild=True,
            )
        except Exception as e:
            return _fail(stage, "VEC-E001",
                         f"Build failed for {cell.mno}/{cell.release}: {e}",
                         time.time() - t0)
        config.save_json(cell_out / "config.json")
        build_stats.save_json(cell_out / "build_stats.json")
        total_chunks += build_stats.total_chunks
        model = build_stats.embedding_model

    stats = {"cells": len(cells), "chunks": total_chunks, "model": model}
    return StageResult(stage=stage, status="OK", elapsed_seconds=time.time() - t0, stats=stats)


# ---------------------------------------------------------------------------
# Stage 9: eval
# ---------------------------------------------------------------------------

def run_eval(ctx: PipelineContext) -> StageResult:
    """Run evaluation."""
    t0 = time.time()
    stage = "eval"
    out_dir = ctx.stage_output(stage)
    out_dir.mkdir(parents=True, exist_ok=True)

    try:
        from core.src.eval.questions import ALL_QUESTIONS
        from core.src.eval.runner import EvalRunner
        from core.src.query.pipeline import build_stub_graph_from_store, load_graph
        from core.src.vectorstore import make_embedder
        from core.src.vectorstore.config import VectorStoreConfig
        from core.src.vectorstore.store_chroma import ChromaDBStore
    except ImportError as e:
        return _fail(stage, "EVL-E001", f"Import error: {e}", time.time() - t0)

    # Load graph; fall back to a stub if absent (RAG-only mode —
    # graph stage was skipped via --rag-only / --skip-graph). The stub
    # has only MNO/Release/Plan nodes so resolver works; pipeline is
    # set to bypass Stage 3 below.
    graph_path = Path(ctx.state.get("graph_path", ctx.stage_output("graph") / "knowledge_graph.json"))
    rag_only_mode = not graph_path.exists()

    # Load vector store
    vs_dir = ctx.stage_output("vectorstore")
    # D-DRAFT-11: per-cell layout has no flat config.json — read the embedder
    # config from a per-cell config (shared across cells) so the embedder
    # matches what built the cells (else defaults to sentence-transformers/HF).
    from core.src.vectorstore.cell_loader import embedder_config_path
    vs_config_path = embedder_config_path(vs_dir)
    vs_config = (
        VectorStoreConfig.load_json(vs_config_path)
        if vs_config_path is not None
        else VectorStoreConfig(
            persist_directory=str(vs_dir),
            embedding_provider=ctx.embedding_provider,
            embedding_model=ctx.embedding_model,
        )
    )

    embedder = make_embedder(vs_config)
    # D-DRAFT-11: load per-cell stores (flat fallback covers a legacy single
    # store). The representative store seeds the rag-only stub graph + back-compat.
    from core.src.vectorstore.cell_loader import load_cell_stores
    cell_stores = load_cell_stores(vs_dir)
    if not cell_stores:
        return _fail(stage, "PIP-E002",
                     f"No vector store under {vs_dir} (run the vectorstore stage first)",
                     time.time() - t0)
    store = next(iter(cell_stores.values()))

    if rag_only_mode:
        graph = build_stub_graph_from_store(store)
    else:
        graph = load_graph(graph_path)

    # Load user-supplied eval questions from Excel if available
    questions = list(ALL_QUESTIONS)
    user_questions = _load_user_eval_questions(ctx.eval_dir)
    if user_questions:
        questions.extend(user_questions)

    # Create synthesizer
    synthesizer = None
    llm = ctx.create_llm_provider(require_real=False)
    if llm and not hasattr(llm, "_is_mock"):
        from core.src.query.synthesizer import LLMSynthesizer
        synthesizer = LLMSynthesizer(llm)

    # Pre-retrieval query rewriter — when a real LLM is available,
    # use it to expand concept-shaped queries with telecom-specific
    # paraphrases before retrieval. Per-query-type policy gates which
    # query types actually call the rewriter (see
    # `pipeline._TYPE_REWRITE_ENABLED`); this construction just makes
    # the rewriter available. Falls back to the mock (no rewrites)
    # when the LLM is unavailable, preserving deterministic behavior.
    rewriter = None
    if llm and not hasattr(llm, "_is_mock"):
        from core.src.query.rewriter import LLMQueryRewriter
        rewriter = LLMQueryRewriter(llm)

    # Cross-encoder reranker — final-pass relevance scoring on the fused
    # top-K. Gated on `resolve_reranker_enabled()` (NORA_RERANKER_ENABLED >
    # DB > config/llm.json > env config > default False) so that, with the
    # reranker disabled (the default), the eval stage never constructs the
    # cross-encoder and never reaches out to HuggingFace for its model —
    # matching the web query path. When enabled, construction is graceful:
    # an uncached/offline model degrades to a MockReranker passthrough.
    from core.src.env.config import resolve_reranker_enabled
    reranker = None
    if resolve_reranker_enabled():
        from core.src.query.reranker import CrossEncoderReranker
        reranker = CrossEncoderReranker()
        if not reranker.available:
            reranker = None  # falls back to MockReranker in pipeline

    runner = EvalRunner(
        graph=graph, embedder=embedder, store=store, cell_stores=cell_stores,
        synthesizer=synthesizer, rewriter=rewriter, reranker=reranker,
    )
    report = runner.run_all(questions, bypass_graph=rag_only_mode)

    # Save report
    report_path = out_dir / "report.json"
    with open(report_path, "w") as f:
        json.dump(report.to_dict(), f, indent=2)

    stats = {
        "questions": len(questions),
        "user_questions": len(user_questions),
        "overall": f"{report.avg_overall:.1%}",
        "completeness": f"{report.avg_completeness:.1%}",
        "accuracy": f"{report.avg_accuracy:.1%}",
        "citation": f"{report.avg_citation_quality:.1%}",
    }
    return StageResult(stage=stage, status="OK", elapsed_seconds=time.time() - t0, stats=stats)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _load_user_eval_questions(eval_dir: Path | None) -> list:
    """Load user-supplied evaluation questions from Excel files.

    Expected Excel columns:
        question_id, category, question, expected_plans (comma-sep),
        expected_req_ids (comma-sep), expected_features (comma-sep),
        expected_standards (comma-sep), expected_concepts (comma-sep),
        min_plans (int), min_chunks (int)
    """
    if not eval_dir or not eval_dir.exists():
        return []

    xlsx_files = list(eval_dir.glob("*.xlsx")) + list(eval_dir.glob("*.xls"))
    if not xlsx_files:
        return []

    questions = []
    try:
        import openpyxl
    except ImportError:
        logger.warning("openpyxl not available — cannot load Excel eval questions")
        return []

    for xlsx_path in xlsx_files:
        try:
            wb = openpyxl.load_workbook(xlsx_path, read_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) < 2:
                continue

            headers = [str(h).strip().lower() if h else "" for h in rows[0]]
            for row in rows[1:]:
                data = {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
                if not data.get("question"):
                    continue

                from core.src.eval.questions import EvalQuestion, GroundTruth
                q = EvalQuestion(
                    question_id=str(data.get("question_id", f"USER_{len(questions)+1:03d}")),
                    category=str(data.get("category", "general")),
                    question=str(data["question"]),
                    ground_truth=GroundTruth(
                        expected_plans=_split_csv(data.get("expected_plans", "")),
                        expected_req_ids=_split_csv(data.get("expected_req_ids", "")),
                        expected_features=_split_csv(data.get("expected_features", "")),
                        expected_standards=_split_csv(data.get("expected_standards", "")),
                        expected_concepts=_split_csv(data.get("expected_concepts", "")),
                        min_plans=int(data.get("min_plans", 1)),
                        min_chunks=int(data.get("min_chunks", 1)),
                    ),
                )
                questions.append(q)
            wb.close()
        except Exception as e:
            logger.warning(f"Failed to load {xlsx_path.name}: {e}")

    if questions:
        logger.info(f"Loaded {len(questions)} user eval questions from {eval_dir}")
    return questions


def _split_csv(value) -> list[str]:
    """Split comma-separated string into list, stripping whitespace."""
    if not value:
        return []
    return [s.strip() for s in str(value).split(",") if s.strip()]


# ---------------------------------------------------------------------------
# Stage function registry
# ---------------------------------------------------------------------------

STAGE_FUNCS: dict[str, callable] = {
    "extract": run_extract,
    "profile": run_profile,
    "parse": run_parse,
    "resolve": run_resolve,
    "taxonomy": run_taxonomy,
    "standards": run_standards,
    "graph": run_graph,
    "vectorstore": run_vectorstore,
    "eval": run_eval,
}
